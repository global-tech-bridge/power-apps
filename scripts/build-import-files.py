#!/usr/bin/env python3
"""SharePoint へ画面から取り込むためのインポート用ファイルを生成する。

    python3 scripts/build-import-files.py

生成物（data/import/）:
  OrgMaster.xlsx       … 「Excel からリストを作成」用。Excel テーブル化してある
  OrgMaster_grid.tsv   … 既存リストへ「グリッド ビュー」で貼り付ける用
  AppAdmins_grid.tsv   … 管理者リストの記入テンプレート

ConsentMaster のインポート用ファイルは意図的に作っていない。
同意文面の本文は 3,000 字を超える複数行テキストで、Excel からリストを作成すると
「1行テキスト（255字）」の列が作られて本文が切り捨てられる。
同意文面は証跡そのものなので、切り捨てが起きる経路を用意しない。
リストの列を先に作ってからフォームで貼り付ける（docs/08-manual-setup.md）。
"""
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC = Path("data/OrgMaster.csv")
DST = Path("data/import")
DST.mkdir(parents=True, exist_ok=True)

with SRC.open(encoding="utf-8-sig", newline="") as f:
    rows = list(csv.DictReader(f))
header = list(rows[0])

# ---------------------------------------------------------------------------
# 1. OrgMaster.xlsx — 「Excel からリストを作成」用
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "OrgMaster"
ws.append(header)
for r in rows:
    ws.append([
        int(r[c]) if c == "SortOrder" else
        (r[c] == "TRUE") if c == "IsActive" else
        r[c]
        for c in header
    ])

# SharePoint の「Excel からリストを作成」は、ファイル内にテーブルがあることを要求する
ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
table = Table(displayName="OrgMaster", ref=ref)
table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
ws.add_table(table)

head_fill = PatternFill("solid", fgColor="DEEBF7")
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center")
widths = {"Title": 12, "BranchName": 14, "BlockCode": 12, "BlockName": 22,
          "SiteCode": 12, "SiteName": 26, "IsActive": 10, "SortOrder": 11}
for i, c in enumerate(header, start=1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
ws.freeze_panes = "A2"

wb.save(DST / "OrgMaster.xlsx")
print(f"data/import/OrgMaster.xlsx: {len(rows)} 行 + ヘッダー（テーブル名 OrgMaster）")

# ---------------------------------------------------------------------------
# 2. OrgMaster_grid.tsv — グリッド ビューへの貼り付け用
#    IsActive（はい/いいえ列）は貼り付けの相性が悪いので列ごと外し、
#    リスト側の既定値を「はい」にしてもらう前提にする。
# ---------------------------------------------------------------------------
grid_cols = [c for c in header if c != "IsActive"]
with (DST / "OrgMaster_grid.tsv").open("w", encoding="utf-8", newline="\n") as f:
    f.write("\t".join(grid_cols) + "\n")
    for r in rows:
        f.write("\t".join(r[c] for c in grid_cols) + "\n")
print(f"data/import/OrgMaster_grid.tsv: {len(rows)} 行（列: {' / '.join(grid_cols)}）")

# ---------------------------------------------------------------------------
# 3. AppAdmins_grid.tsv — 記入テンプレート
# ---------------------------------------------------------------------------
with (DST / "AppAdmins_grid.tsv").open("w", encoding="utf-8", newline="\n") as f:
    f.write("Title\tUserEmail\tUserName\tNote\n")
    f.write("taro.yamada@example.co.jp\ttaro.yamada@example.co.jp\t山田 太郎\t2026-09-09 ○○課長承認\n")
print("data/import/AppAdmins_grid.tsv: 記入例1行（実際のメールアドレスに書き換えて使う）")
